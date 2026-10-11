"""Opt-in real Collabora CODE test, not a simulated office iframe."""
from io import BytesIO
import os
import socket
import threading
import time

import pytest
import uvicorn
from django.contrib.staticfiles.handlers import ASGIStaticFilesHandler
from docx import Document as WordDocument
from playwright.sync_api import expect, sync_playwright

from wiki.services import save_document
from wiki.storage import active_storage
from tests.browser.test_word_editor import login

pytestmark = [pytest.mark.browser, pytest.mark.django_db(transaction=True),
              pytest.mark.skipif(os.environ.get('RUN_COLLABORA') != '1', reason='Requires real Collabora at port 19980')]


@pytest.fixture
def office_server(settings, workspace):
    from config.asgi import application
    settings.PUBLIC_URL = 'http://127.0.0.1:18080'
    settings.COLLABORA_URL = settings.COLLABORA_INTERNAL_URL = 'http://127.0.0.1:19980/collabora'
    settings.COLLABORA_WOPI_URL = 'http://host.docker.internal:18080'
    settings.ALLOWED_HOSTS = ['127.0.0.1', 'host.docker.internal', 'testserver']
    sock = socket.socket(); sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1); sock.bind(('0.0.0.0', 18080))
    server = uvicorn.Server(uvicorn.Config(ASGIStaticFilesHandler(application), log_level='warning', lifespan='on'))
    thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True); thread.start()
    for _ in range(100):
        if server.started: break
        time.sleep(.05)
    assert server.started
    yield settings.PUBLIC_URL
    server.should_exit = True; thread.join(timeout=10); sock.close()


def test_real_docx_edit_and_save(office_server, workspace):
    doc = save_document(workspace['admin'], 'Collabora real document', 'document', 'Notes', 'Initial office text', workspace['team'])
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=os.environ.get('PLAYWRIGHT_CHROMIUM_EXECUTABLE_PATH'))
        page = browser.new_page(viewport={'width': 1440, 'height': 1000})
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        login(page, office_server, 'admin')
        page.goto(office_server + f'/documents/{doc.pk}/live/')
        page.wait_for_url('**/office/')
        try:
            expect(page.locator('#office-status')).to_contain_text('Office editor ready', timeout=60000)
        finally:
            page.screenshot(path='/tmp/yourwiki-collabora-loading.png')
        frame = page.frame_locator('#office-frame')
        expect(frame.locator('#document-container')).to_be_visible()
        welcome = frame.locator('.iframe-welcome-wrap')
        if welcome.count():
            intro = frame.frame_locator('.iframe-welcome-wrap iframe')
            for step in (1, 2, 3):
                intro.locator(f'#slide-{step}-button').click(timeout=5000)
        page.screenshot(path='/tmp/yourwiki-collabora-editor.png')
        frame.locator('#document-container').click(position={'x': 350, 'y': 200})
        page.keyboard.press('Control+End')
        page.keyboard.press('Enter')
        page.keyboard.type('Saved through WOPI')
        page.keyboard.press('Control+s')
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            # Fetch from a worker because Playwright's sync loop runs an asyncio loop.
            from concurrent.futures import ThreadPoolExecutor
            def contents():
                from wiki.models import Document
                current = Document.objects.get(pk=doc.pk)
                return '\n'.join(p.text for p in WordDocument(BytesIO(active_storage().read(current.reference))).paragraphs)
            with ThreadPoolExecutor(max_workers=1) as pool:
                text = pool.submit(contents).result()
            if 'Saved through WOPI' in text: break
            page.wait_for_timeout(250)
        assert 'Saved through WOPI' in text, text
        assert not errors, errors
        browser.close()


@pytest.mark.parametrize('format', ['pptx', 'md', 'xlsx', 'ods'])
def test_real_additional_office_formats_load(office_server, workspace, format):
    if format == 'pptx':
        doc = save_document(workspace['admin'], 'Office slides', 'presentation', 'Notes', 'Slide content', workspace['team'])
    elif format == 'md':
        doc = save_document(workspace['admin'], 'Office markdown', 'document', 'Notes', '# Markdown', workspace['team'], file_format='md')
    elif format == 'ods':
        doc = save_document(workspace['admin'], 'Office spreadsheet', 'table', 'Notes', 'Name,Value\nTest,42', workspace['team'])
    else:
        import xlsxwriter
        output = BytesIO(); book = xlsxwriter.Workbook(output); book.add_worksheet().write('A1', 'Excel'); book.close()
        doc = save_document(workspace['admin'], 'Office.xlsx', 'file', 'Notes', output.getvalue(), workspace['team'])
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=os.environ.get('PLAYWRIGHT_CHROMIUM_EXECUTABLE_PATH'))
        page = browser.new_page(viewport={'width':1440, 'height':1000})
        login(page, office_server, 'admin')
        page.goto(office_server + f'/documents/{doc.pk}/live/')
        page.wait_for_url('**/office/')
        try:
            expect(page.locator('#office-status')).to_contain_text('Office editor ready', timeout=60000)
            expect(page.frame_locator('#office-frame').locator('#document-container')).to_be_visible()
        finally:
            page.screenshot(path=f'/tmp/yourwiki-collabora-{format}.png')
            browser.close()


def test_real_collabora_shared_cursor_and_readonly(office_server, workspace):
    doc = save_document(workspace['admin'], 'Shared office document', 'document', 'Notes',
                        'Shared cursor position', workspace['team'])
    doc.inherit_permissions = False
    doc.policy['group'] = {'read': True, 'write': True}
    doc.policy['everyone'] = {'read': True, 'write': False}
    doc.save()
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=os.environ.get('PLAYWRIGHT_CHROMIUM_EXECUTABLE_PATH'))
        pages = []
        for user in ('admin', 'member', 'outsider'):
            page = browser.new_context(viewport={'width': 1440, 'height': 1000}).new_page()
            login(page, office_server, user)
            page.goto(office_server + f'/documents/{doc.pk}/office/')
            expect(page.locator('#office-status')).to_contain_text('Office editor ready', timeout=60000)
            frame = page.frame_locator('#office-frame')
            if frame.locator('.iframe-welcome-wrap').count():
                intro = frame.frame_locator('.iframe-welcome-wrap iframe')
                for step in (1, 2, 3):
                    intro.locator(f'#slide-{step}-button').click(timeout=5000)
            pages.append(page)
        a, b, reader = pages
        bf = b.frame_locator('#office-frame')
        bf.locator('#document-container').click(position={'x': 350, 'y': 200})
        b.keyboard.press('Control+End')
        b.keyboard.press('ArrowLeft')
        expect(a.frame_locator('#office-frame').locator('.cursor-header-section').filter(has_text='Member')).to_be_visible(timeout=15000)
        expect(reader.locator('#office-status')).to_contain_text('View only')
        rf = next(f for f in reader.frames if '/browser/' in f.url)
        assert rf.evaluate('() => app.isReadOnly()') is True
        reader.screenshot(path='/tmp/yourwiki-collabora-readonly.png')
        a.screenshot(path='/tmp/yourwiki-collabora-shared-cursor.png')
        browser.close()


def test_real_collabora_offers_open_fonts(office_server, workspace):
    doc = save_document(workspace['admin'], 'Collabora fonts', 'document', 'Notes', 'Fonts', workspace['team'])
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=os.environ.get('PLAYWRIGHT_CHROMIUM_EXECUTABLE_PATH'))
        page = browser.new_page(viewport={'width':1440,'height':1000})
        login(page, office_server, 'admin')
        page.goto(office_server+f'/documents/{doc.pk}/office/')
        expect(page.locator('#office-status')).to_contain_text('Office editor ready',timeout=60000)
        frame = page.frame(name='office-frame') or next(f for f in page.frames if '/browser/' in f.url)
        values = frame.evaluate("() => JSON.stringify(window.app?.map?._docLayer?._toolbarCommandValues || {})")
        assert 'Tinos' in values and 'Oswald' in values, values[:2000]
        assert 'Anton' in values and 'Roboto' in values
        if os.environ.get('REQUIRE_IMPACT') == '1':
            assert 'Impact' in values
        browser.close()
