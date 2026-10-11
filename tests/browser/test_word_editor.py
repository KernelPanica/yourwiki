"""DOCX/Markdown visual editing, real collaboration and portable output."""
import io
import os
import re
from pathlib import Path
import pytest
from playwright.sync_api import expect, sync_playwright
from wiki.services import save_document

pytestmark = [pytest.mark.browser, pytest.mark.django_db(transaction=True),
              pytest.mark.skipif(os.environ.get('RUN_BROWSER') != '1', reason='Set RUN_BROWSER=1')]


def login(page, server, username):
    page.goto(server + '/login/')
    page.get_by_label('Username').fill(username)
    page.get_by_label('Password', exact=True).fill('Correct-password-8923')
    page.get_by_role('button', name='Sign in', exact=True).click()


@pytest.mark.parametrize('file_format', ['docx', 'md'])
def test_word_gui_edits_merges_and_exports(realtime_server, workspace, file_format):
    doc = save_document(workspace['admin'], 'Office document', 'document', 'Notes',
                        'Original paragraph', workspace['team'], file_format=file_format)
    doc.inherit_permissions = False
    doc.policy['group']['write'] = True
    doc.save()
    errors = []
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=os.environ.get('PLAYWRIGHT_CHROMIUM_EXECUTABLE_PATH'))
        a = browser.new_page(viewport={'width':1440,'height':1000})
        b = browser.new_page(viewport={'width':1440,'height':1000})
        for page, username in [(a, 'admin'), (b, 'member')]:
            page.on('pageerror', lambda error: errors.append(str(error)))
            login(page, realtime_server, username)
            page.goto(realtime_server + f'/documents/{doc.pk}/live/')
            expect(page.locator('.document-host .tiptap')).to_contain_text('Original paragraph')
            expect(page.locator('.sidebar')).to_have_attribute('aria-hidden', 'true')
            page.locator('[data-sidebar-open]').click()
            expect(page.locator('.sidebar')).to_have_attribute('aria-hidden', 'false')
            page.locator('[data-sidebar-open]').click()
            expect(page.locator('.topbar')).to_be_visible()
            expect(page.locator('body')).not_to_have_class(re.compile('docs-page'))
            expect(page.locator('main .docs-page #visual-editor')).to_be_visible()
            expect(page.get_by_role('group', name='Formatting', exact=True)).to_be_visible()
            page.locator('.docs-menus summary').filter(has_text=re.compile('^Insert$')).click()
            expect(page.get_by_role('group', name='Insert', exact=True)).to_be_visible()
            page.keyboard.press('Escape')
            expect(page.locator('#editor-error')).to_be_hidden()
            expect(page.locator('#document-actions #save-live')).to_be_visible()
            expect(page.locator('.docs-menus summary').filter(has_text=re.compile('^File$'))).to_have_count(0)
            expect(page.locator('.docs-page #review-panel')).to_have_count(0)
        a.bring_to_front()
        a.locator('.tiptap').click()
        a.keyboard.press('Control+End'); a.keyboard.type(' Alice')
        a.keyboard.press('Control+Shift+ArrowLeft')
        a.get_by_role('button', name='Bold', exact=True).click()
        expect(b.locator('.tiptap strong')).to_contain_text('Alice', timeout=15000)
        b.bring_to_front()
        b.locator('.tiptap').click()
        b.keyboard.press('Control+End'); b.keyboard.type(' Bob')
        expect(a.locator('.tiptap')).to_contain_text('Bob', timeout=15000)
        trace={'after_bob':a.locator('.tiptap').inner_html()}
        a.bring_to_front()
        a.locator('.tiptap').click()
        a.keyboard.press('Control+End'); a.keyboard.press('Enter')
        trace['after_enter']=a.locator('.tiptap').evaluate('(el)=>({html:el.innerHTML,selection:el.editor?.state.selection.toJSON()})')
        a.locator('.docs-menus summary').filter(has_text=re.compile('^Insert$')).click()
        trace['before_table']=a.locator('.tiptap').evaluate('(el)=>({html:el.innerHTML,selection:el.editor?.state.selection.toJSON()})')
        a.get_by_role('button', name='Table', exact=True).click()
        trace['after_table']=a.locator('.tiptap').inner_html()
        expect(b.locator('.tiptap table')).to_be_visible(timeout=15000)
        a.get_by_role('button', name='Save now', exact=True).click()
        expect(a.locator('#sync-status')).to_have_text(re.compile('^(Saved|Synced) to storage$'), timeout=15000)
        export = a.request.get(realtime_server + f'/documents/{doc.pk}/export/')
        assert export.status == 200
        if file_format == 'docx':
            from docx import Document as Word
            word = Word(io.BytesIO(export.body()))
            assert any('Alice' in paragraph.text and 'Bob' in paragraph.text for paragraph in word.paragraphs), {'paragraphs':[paragraph.text for paragraph in word.paragraphs], 'trace':trace}
            assert any(run.bold and 'Alice' in run.text for paragraph in word.paragraphs for run in paragraph.runs)
            assert len(word.tables) == 1
        a.reload()
        expect(a.locator('.tiptap')).to_contain_text('Bob')
        expect(a.locator('.tiptap table')).to_be_visible()
        a.locator('.docs-menus summary').filter(has_text=re.compile('^Edit$')).focus()
        a.keyboard.press('ArrowDown')
        expect(a.locator('.docs-menus details').first).to_have_attribute('open', '')
        a.keyboard.press('Escape')
        expect(a.locator('.docs-menus details').first).not_to_have_attribute('open', '')
        a.locator('#document-actions').get_by_role('button', name='Comments and suggestions', exact=True).click()
        expect(a.locator('#review-panel')).to_be_visible()
        a.locator('#document-actions').get_by_role('button', name='Comments and suggestions', exact=True).click()
        expect(a.locator('#review-panel')).to_be_hidden()
        a.get_by_label('Zoom', exact=True).select_option('75')
        assert a.locator('.tiptap').evaluate('e => e.style.zoom') == '0.75'
        a.get_by_label('Zoom', exact=True).select_option('100')
        a.locator('[data-editor-fullscreen]').click()
        expect(a.locator('body')).to_have_class(re.compile('editor-fullscreen'))
        expect(a.locator('.tiptap')).to_contain_text('Bob')
        a.locator('[data-editor-fullscreen]').click()
        expect(a.locator('body')).not_to_have_class(re.compile('editor-fullscreen'))
        output = Path('/tmp/yourwiki-word-editor'); output.mkdir(exist_ok=True)
        a.screenshot(path=str(output/f'{file_format}-desktop.png'), full_page=True)
        a.context.add_cookies([{'name':'django_language','value':'ru','url':realtime_server}])
        a.reload()
        expect(a.locator('.tiptap')).to_contain_text('Bob')
        a.screenshot(path=str(output/f'{file_format}-desktop-ru.png'), full_page=True)
        a.set_viewport_size({'width':320,'height':900})
        assert a.evaluate('document.documentElement.scrollWidth <= innerWidth + 1')
        assert a.locator('#editor-toolbar').evaluate('e => e.clientWidth <= innerWidth')
        a.screenshot(path=str(output/f'{file_format}-mobile.png'), full_page=True)
        assert not errors, errors
        browser.close()
