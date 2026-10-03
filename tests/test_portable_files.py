import base64
import io
import json
import zipfile
from unittest.mock import patch
import pytest
from django.core.exceptions import ValidationError
from wiki.models import Document, Folder, SiteConfiguration, Collaboration
from wiki.services import save_document, Conflict
from wiki.storage import active_storage, GoogleDrive
from wiki.file_formats import encode, decode, read_ods, read_docx
from wiki.native import pack, unpack
from wiki.collaboration import ensure_room, current_content

pytestmark = pytest.mark.django_db


def test_history_inheritance_stale_save_and_physical_rename(workspace, document):
    user, team = workspace['admin'], workspace['team']
    root = Folder.objects.create(name='Root', owner=user, group=team, revision_history=True)
    child = Folder.objects.create(name='Child', parent=root, owner=user, group=team)
    document.folder = child
    document.save()
    stale = Document.objects.get(pk=document.pk)
    saved = save_document(user, 'Renamed', 'document', 'Notes', 'second', team, document, document.revision)
    assert not (workspace['root'] / stale.reference).exists()
    assert saved.reference == 'Root/Child/Renamed.md'
    history = json.loads(active_storage().read(saved.history_reference))
    assert base64.b64decode(history['entries'][0]['content_base64']).startswith(b'# Hello')
    with pytest.raises(Conflict):
        save_document(user, 'Stale', 'document', 'Notes', 'lost update', team, stale, stale.revision)
    assert active_storage().read(saved.reference) == b'second'
    child.revision_history = False
    child.save()
    saved = save_document(user, saved.title, 'document', 'Notes', 'third', team, saved, saved.revision)
    assert len(json.loads(active_storage().read(saved.history_reference))['entries']) == 1
    saved.revision_history = True
    saved.save()
    saved = save_document(user, saved.title, 'document', 'Notes', 'fourth', team, saved, saved.revision)
    assert len(json.loads(active_storage().read(saved.history_reference))['entries']) == 2


def test_portable_files_restore_editor_without_database_state(workspace):
    rich = pack('document', {'type': 'doc', 'content': [
        {'type':'heading', 'attrs': {'level': 2}, 'content':[{'type':'text','text':'Привет','marks':[{'type':'bold'}]}]},
        {'type':'bulletList','content':[{'type':'listItem','content':[{'type':'paragraph','content':[{'type':'text','text':'list'}]}]}]},
        {'type':'table','content':[{'type':'tableRow','content':[{'type':'tableCell','content':[{'type':'paragraph','content':[{'type':'text','text':'cell'}]}]}]}]}
    ]})
    for fmt in ('md','docx'):
        doc = save_document(workspace['admin'], fmt, 'document', 'Notes', rich, workspace['team'], file_format=fmt)
        assert doc.reference.endswith('.'+fmt)
        assert not doc.history_reference
        raw = active_storage().read(doc.reference)
        assert decode(doc, raw) == rich
        room = ensure_room(doc)
        room.delete()
        assert current_content(doc) == rich
        if fmt == 'docx':
            from docx import Document as Word
            word = Word(io.BytesIO(raw))
            assert word.tables[0].cell(0,0).text == 'cell'
            assert 'list' in [p.text for p in word.paragraphs]
    assert not list(workspace['root'].rglob('*.json'))


def test_ods_formula_styles_sheets_and_foreign_import(workspace):
    cells = {'0': {'0': {'v': 2, 's':'bold'}, '1': {'v': 3}}, '1': {'0': {'v':5, 'f':'=SUM(A1:B1)'}}}
    content = pack('table', {'id':'book','sheetOrder':['one','two'], 'styles':{'bold': {'bl':1, 'bg': {'rgb':'#ff0000'}}},
        'sheets':{sid: {'id':sid, 'name':sid, 'rowCount':200, 'columnCount':26, 'cellData':cells} for sid in ('one','two')}})
    doc = save_document(workspace['admin'], 'Sheet', 'table', 'Notes', content, workspace['team'])
    raw = active_storage().read(doc.reference)
    assert doc.reference.endswith('.ods') and read_ods(raw) == content
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(output,'w') as target:
        assert source.read('mimetype') == b'application/vnd.oasis.opendocument.spreadsheet'
        assert b'of:=SUM([.A1:.B1])' in source.read('content.xml')
        assert b'font-weight="bold"' in source.read('content.xml')
        for entry in source.infolist():
            if entry.filename != 'meta.xml': target.writestr(entry, source.read(entry.filename))
    imported = unpack(read_ods(output.getvalue()))['data']
    assert len(imported['sheets']) == 2
    assert imported['sheets']['sheet1']['cellData']['1']['0']['f'] == '=SUM(A1:B1)'
    for reader in (read_ods, read_docx):
        with pytest.raises(ValidationError): reader(b'not a zip archive')


def test_external_conflict_and_explicit_reload(client, workspace, document):
    ensure_room(document)
    active_storage().write(document.reference, b'Changed outside', document.reference)
    with pytest.raises(Conflict):
        save_document(workspace['admin'], document.title, 'document', 'Notes', 'overwrite', workspace['team'], document, document.revision)
    client.force_login(workspace['admin'])
    response = client.post(f'/documents/{document.pk}/file-settings/', {'action':'reload','discard':'on'})
    assert response.status_code == 302
    document.refresh_from_db()
    assert current_content(document) == 'Changed outside'
    assert not Collaboration.objects.filter(document=document).exists()


def test_recursive_acl_owner_and_group_overrides(workspace):
    user, member, team = workspace['admin'], workspace['member'], workspace['team']
    root = Folder.objects.create(name='Root', owner=user, group=team)
    child = Folder.objects.create(name='Child', parent=root, owner=member, group=team)
    doc = save_document(member, 'Inherited', 'document', 'Notes', 'text', team, folder=child)
    assert not doc.allows(member, 'write')  # Creating a child does not grant ownership of inherited ACLs.
    root.group_policies = {str(team.pk): {'visible':True,'read':True,'write':True}}
    root.save()
    doc = Document.objects.get(pk=doc.pk)
    assert doc.allows(member, 'write')
    doc.inherit_permissions = False
    doc.save()
    root.group_policies[str(team.pk)]['write'] = False
    root.save()
    assert Document.objects.get(pk=doc.pk).allows(member, 'write')


def test_original_link_read_permissions_and_provider_move(client, workspace, document):
    client.force_login(workspace['member'])
    with patch('wiki.storage.MountedStorage.original_url', return_value='https://drive.google.com/file/d/private/view'):
        page = client.get(f'/documents/{document.pk}/')
        assert page.status_code == 200 and b'Open original' in page.content
    client.force_login(workspace['outsider'])
    with patch('wiki.storage.MountedStorage.original_url') as url:
        assert client.get(f'/documents/{document.pk}/').status_code != 200
        url.assert_not_called()
    google = GoogleDrive({'root':'root'})
    with patch.object(google, 'headers', return_value={}), patch.object(google, 'directory', return_value='destination'), patch.object(google, 'request') as request:
        request.return_value.json.return_value = {'parents':['root']}
        assert google.move('original-id', 'Folder/file.ods') == 'original-id'
        assert [call.args[0] for call in request.call_args_list] == ['GET', 'PATCH']
        assert request.call_args.kwargs['params']['removeParents'] == 'root'
        assert request.call_args.kwargs['json'] == {'name':'file.ods'}


def test_history_settings_require_write(client, workspace, document):
    client.force_login(workspace['member'])
    assert client.post(f'/documents/{document.pk}/file-settings/', {'revision_history':'on'}).status_code != 302
    client.force_login(workspace['admin'])
    assert client.get(f'/documents/{document.pk}/file-settings/').status_code == 200
    assert client.post(f'/documents/{document.pk}/file-settings/', {'revision_history':'on'}).status_code == 302
    document.refresh_from_db()
    assert document.history_enabled()


def test_failed_first_history_save_is_retryable(workspace, document):
    from wiki.storage import StorageError
    document.revision_history = True
    document.save()
    adapter = active_storage()
    original = adapter.read(document.reference)
    write = adapter.write
    def fail_main(key, data, reference=None, **kwargs):
        if key.endswith('.md'): raise StorageError('Offline')
        return write(key, data, reference, **kwargs)
    from wiki.file_history import persist
    with patch.object(adapter, 'write', side_effect=fail_main):
        with pytest.raises(StorageError): persist(document, 'new', document.reference, adapter=adapter)
    assert not list(workspace['root'].rglob('*.json'))
    assert adapter.read(document.reference) == original
    saved = save_document(workspace['admin'], document.title, 'document', 'Notes', 'new', workspace['team'], document, document.revision)
    assert len(json.loads(adapter.read(saved.history_reference))['entries']) == 1


def test_history_can_hold_a_full_size_file(workspace):
    from wiki.storage import HISTORY_MAX_BYTES
    doc = save_document(workspace['admin'], 'large.bin', 'file', 'Notes', b'x'*(5*1024*1024), workspace['team'])
    doc.revision_history = True
    doc.save()
    doc = save_document(workspace['admin'], doc.title, 'file', 'Notes', b'new', workspace['team'], doc, doc.revision)
    history = json.loads(active_storage().read(doc.history_reference, limit=HISTORY_MAX_BYTES))
    assert len(base64.b64decode(history['entries'][0]['content_base64'])) == 5*1024*1024


def test_legacy_conversion_preserves_native_content_and_is_repeatable(workspace):
    from django.core.management import call_command
    content = pack('document', {'type':'doc','content':[{'type':'paragraph','content':[{'type':'text','text':'Legacy rich content'}]}]})
    ref = active_storage().write('legacy.wiki.json', content.encode())
    doc = Document.objects.create(title='Legacy', kind='document', owner=workspace['admin'], group=workspace['team'], reference=ref, format_version=1)
    call_command('convert_file_formats')
    doc.refresh_from_db()
    assert doc.reference == 'Legacy.md' and doc.file_format == 'md'
    assert current_content(doc) == content
    assert not (workspace['root']/ref).exists()
    call_command('convert_file_formats')
    assert len(list(workspace['root'].iterdir())) == 1


def test_scan_does_not_publish_history_as_an_independent_file(workspace):
    from django.core.management import call_command
    adapter = active_storage()
    adapter.write('private.bin.json', json.dumps({'format':'yourwiki-history','version':1,'entries':[]}).encode())
    call_command('scan_storage')
    assert not Document.objects.exists()


def test_text_document_cannot_claim_a_spreadsheet_format(workspace):
    with pytest.raises(ValidationError):
        save_document(workspace['admin'], 'Wrong format', 'document', 'Notes', 'content', workspace['team'], file_format='ods')
    assert not list(workspace['root'].iterdir())
