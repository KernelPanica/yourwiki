"""WOPI preserves complete office files and enforces tokens, leases and ACL."""
from datetime import timedelta
from io import BytesIO
from unittest.mock import Mock, patch
from urllib.parse import urlencode

import pytest
from django.core import signing
from django.test import Client
from django.utils import timezone
from docx import Document as WordDocument

from wiki.collabora import open_session, token_for, SALT, editor_action
from wiki.models import OfficeSession
from wiki.services import save_document, Conflict
from wiki.storage import active_storage, StorageError

pytestmark = pytest.mark.django_db


@pytest.fixture
def office(workspace, settings):
    settings.COLLABORA_URL = 'https://office.example.test/collabora'
    settings.COLLABORA_INTERNAL_URL = 'http://office:9980/collabora'
    settings.COLLABORA_WOPI_URL = 'https://yourwiki.example.test'
    doc = save_document(workspace['admin'], 'Office file', 'document', 'Notes', 'Original', workspace['team'])
    doc.inherit_permissions = False
    doc.policy['group'] = {'read': True, 'write': True}
    doc.save()
    session = open_session(doc, workspace['admin'])
    return doc, session


def url(session, user, contents=False):
    return f'/wopi/files/{session.pk}' + ('/contents' if contents else '') + '?' + urlencode({'access_token': token_for(session, user)})


def rich_docx():
    doc = WordDocument()
    doc.add_paragraph('Changed in Collabora')
    doc.sections[0].header.paragraphs[0].text = 'Preserve full header'
    doc.add_paragraph('Superscript').runs[0].font.superscript = True
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_shared_sessions_and_exact_binary_save(office, workspace):
    doc, session = office
    assert open_session(doc, workspace['member']).pk == session.pk
    client = Client(enforce_csrf_checks=True)
    info = client.get(url(session, workspace['member']))
    assert info.status_code == 200
    assert info.json()['UserCanWrite'] is True
    assert info.json()['BaseFileName'].endswith('.docx')
    original = client.get(url(session, workspace['member'], True))
    assert original.content == active_storage().read(doc.reference)
    body = rich_docx()
    response = client.post(url(session, workspace['member'], True), body,
                           content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT')
    assert response.status_code == 200
    doc.refresh_from_db()
    assert active_storage().read(doc.reference) == body
    assert response['X-WOPI-ItemVersion'] == doc.storage_digest
    assert client.get(url(session, workspace['admin'])).json()['Version'] == doc.storage_digest
    assert OfficeSession.objects.get(pk=session.pk).pending == b''


def test_tokens_are_bound_expired_and_revoked(office, workspace):
    doc, session = office
    client = Client()
    assert client.get(f'/wopi/files/{session.pk}').status_code == 403
    assert client.get(url(session, workspace['outsider'])).status_code == 403
    forged = signing.dumps({'session': str(session.pk), 'user': workspace['admin'].pk, 'write': True}, salt='wrong')
    assert client.get(f'/wopi/files/{session.pk}?access_token={forged}').status_code == 403
    from uuid import uuid4
    assert client.get(url(session, workspace['admin']).replace(str(session.pk), str(uuid4()), 1)).status_code == 403
    workspace['member'].is_active = False
    workspace['member'].save()
    assert client.get(url(session, workspace['member'])).status_code == 403
    OfficeSession.objects.filter(pk=session.pk).update(expires_at=timezone.now()-timedelta(seconds=1))
    assert client.get(url(session, workspace['admin'])).status_code == 403


def test_readonly_token_and_revoked_write_cannot_save(office, workspace):
    doc, session = office
    doc.policy['group'] = {'read': True, 'write': False}
    doc.save()
    session.document = doc
    readonly = url(session, workspace['member'], True)
    assert Client().post(readonly, rich_docx(), content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT').status_code == 403
    assert Client().get(url(session, workspace['member'])).json()['UserCanWrite'] is False
    doc.policy['group']['write'] = True
    doc.save()
    # A read-only capability cannot acquire write access through an ACL change.
    assert Client().post(readonly, rich_docx(), content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT').status_code == 403
    writable = url(session, workspace['member'], True)
    doc.policy['group']['write'] = False
    doc.save()
    assert Client().post(writable, rich_docx(), content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT').status_code == 403


def test_lock_mismatch_and_refresh(office, workspace):
    _, session = office
    client = Client()
    endpoint = url(session, workspace['admin'])
    assert client.post(endpoint, HTTP_X_WOPI_OVERRIDE='LOCK', HTTP_X_WOPI_LOCK='one').status_code == 200
    mismatch = client.post(endpoint, HTTP_X_WOPI_OVERRIDE='LOCK', HTTP_X_WOPI_LOCK='two')
    assert mismatch.status_code == 409 and mismatch['X-WOPI-Lock'] == 'one'
    assert client.post(endpoint, HTTP_X_WOPI_OVERRIDE='REFRESH_LOCK', HTTP_X_WOPI_LOCK='one').status_code == 200
    OfficeSession.objects.filter(pk=session.pk).update(pending=b'recovery from valid editor', pending_author=workspace['admin'])
    bad_save = client.post(url(session, workspace['admin'], True), rich_docx(), content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT', HTTP_X_WOPI_LOCK='two')
    assert bad_save.status_code == 409
    session.refresh_from_db()
    assert bytes(session.pending) == b'recovery from valid editor'
    assert client.post(endpoint, HTTP_X_WOPI_OVERRIDE='UNLOCK', HTTP_X_WOPI_LOCK='one').status_code == 200


def test_external_conflict_keeps_recovery_and_original(office, workspace):
    doc, session = office
    external = rich_docx()
    active_storage().write('Notes/Office file.docx', external, doc.reference)
    body = BytesIO()
    changed = WordDocument(); changed.add_paragraph('Our separate edits'); changed.save(body)
    client = Client()
    response = client.post(url(session, workspace['admin'], True), body.getvalue(), content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT')
    assert response.status_code == 409
    assert active_storage().read(doc.reference) == external
    assert bytes(OfficeSession.objects.get(pk=session.pk).pending) == body.getvalue()
    client.force_login(workspace['admin'])
    assert client.get(f'/documents/{doc.pk}/office/recovery/').content == body.getvalue()
    client.force_login(workspace['member'])
    assert client.get(f'/documents/{doc.pk}/office/recovery/').status_code != 200


def test_other_editors_are_blocked_and_old_sessions_rotate(office, workspace):
    doc, session = office
    from wiki.source import check_available, acquire
    from wiki.saves import write_snapshot
    for action in [lambda: check_available(doc), lambda: acquire(doc, workspace['admin']),
                   lambda: write_snapshot(doc, 'No overwrite', 'Notes/Office file.docx', user=workspace['admin'])]:
        with pytest.raises(Conflict):
            action()
    previous = url(session, workspace['admin'])
    OfficeSession.objects.filter(pk=session.pk).update(expires_at=timezone.now()-timedelta(seconds=1))
    assert open_session(doc, workspace['admin']).pk != session.pk
    assert Client().get(previous).status_code == 403


def test_discovery_ui_has_no_provider_secrets(office, workspace):
    doc, _ = office
    xml = b'<wopi-discovery><net-zone><app><action name="edit" ext="docx" urlsrc="http://internal/collabora/browser/build/cool.html?"/></app></net-zone></wopi-discovery>'
    with patch('wiki.collabora.requests.get', return_value=Mock(content=xml)):
        client = Client(); client.force_login(workspace['admin'])
        response = client.get(f'/documents/{doc.pk}/office/')
        assert response.status_code == 200
        assert b'https://office.example.test/collabora/browser/build/cool.html?' in response.content
        assert 'https://office.example.test' in response['Content-Security-Policy']
        assert b'office_token' not in response.content
    assert Client().get('/wopi/files/not-an-id').status_code == 404


def test_storage_failure_retains_bytes_for_retry(office, workspace):
    doc, session = office
    body = rich_docx()
    with patch('wiki.saves.write_snapshot', side_effect=StorageError('unavailable')):
        response = Client().post(url(session, workspace['admin'], True), body, content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT')
        assert response.status_code == 503
    assert bytes(OfficeSession.objects.get(pk=session.pk).pending) == body
    assert Client().post(url(session, workspace['admin'], True), body, content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT').status_code == 200


def test_invalid_upload_and_ods_binary_save(office, workspace):
    _, session = office
    assert Client().post(url(session, workspace['admin'], True), b'broken', content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT').status_code == 400
    doc = save_document(workspace['admin'], 'Office sheet', 'table', 'Notes', '10\n20', workspace['team'])
    session = open_session(doc, workspace['admin'])
    data = active_storage().read(doc.reference)
    assert Client().post(url(session, workspace['admin'], True), data, content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT').status_code == 200
    doc.refresh_from_db()
    assert active_storage().read(doc.reference) == data


def test_legacy_docx_relationships_are_repaired_without_storage_write(office, workspace):
    import zipfile
    from xml.etree import ElementTree
    from wiki.collabora import editor_bytes
    doc, session = office
    data = active_storage().read(doc.reference)
    output = BytesIO()
    part = 'word/_rels/document.xml.rels'
    with zipfile.ZipFile(BytesIO(data)) as source, zipfile.ZipFile(output, 'w') as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == part:
                content = ElementTree.tostring(ElementTree.fromstring(content))
                assert b'ns0:Relationships' in content
            target.writestr(entry, content)
    legacy = output.getvalue()
    repaired = editor_bytes(doc, legacy)
    with zipfile.ZipFile(BytesIO(legacy)) as before, zipfile.ZipFile(BytesIO(repaired)) as after:
        assert b'ns0:Relationships' not in after.read(part)
        for name in before.namelist():
            if name != part:
                assert before.read(name) == after.read(name)
    assert editor_bytes(doc, repaired) == repaired
    assert active_storage().read(doc.reference) == data


def test_expired_failed_session_reopens_with_recovery(office, workspace):
    doc, session = office
    OfficeSession.objects.filter(pk=session.pk).update(
        expires_at=timezone.now()-timedelta(minutes=1), pending=rich_docx(),
        pending_author=workspace['admin'], error='Saving failed.')
    reopened = open_session(doc, workspace['admin'])
    assert reopened.pk == session.pk
    assert reopened.expires_at > timezone.now()
    assert reopened.pending
    assert Client().get(url(reopened, workspace['admin'])).status_code == 200


def test_create_presentation_and_save_complete_pptx(office, workspace):
    from pptx import Presentation
    from wiki.models import Document
    client = Client()
    client.force_login(workspace['admin'])
    response = client.post('/documents/new/', {'title':'Slides.pptx', 'kind':'presentation', 'content':'First slide'})
    assert response.status_code == 302
    doc = Document.objects.get(kind='presentation')
    assert doc.title == 'Slides' and doc.file_format == 'pptx'
    original = active_storage().read(doc.reference)
    deck = Presentation(BytesIO(original))
    assert len(deck.slides) == 1
    assert deck.slides[0].shapes.title.text == 'Slides'
    session = open_session(doc, workspace['admin'])
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = 'Second slide'
    slide.notes_slide.notes_text_frame.text = 'Keep speaker notes'
    output = BytesIO(); deck.save(output)
    result = client.post(url(session, workspace['admin'], True), output.getvalue(),
                         content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT')
    assert result.status_code == 200
    assert active_storage().read(doc.reference) == output.getvalue()
    assert client.get(f'/documents/{doc.pk}/').status_code == 200
    assert client.get(f'/documents/{doc.pk}/export/').content == output.getvalue()


@pytest.mark.parametrize('name', ['Exact.docx', 'Exact.pptx', 'Exact.xlsx', 'Exact.odt', 'Exact.ods', 'Exact.md'])
def test_uploaded_office_files_remain_exact_and_open_in_collabora(office, workspace, name):
    from django.core.files.uploadedfile import SimpleUploadedFile
    from wiki.models import Document
    if name.endswith('.docx'):
        data = rich_docx()
    elif name.endswith('.pptx'):
        from wiki.office_formats import presentation
        data = presentation('Uploaded')
    elif name.endswith('.xlsx'):
        import xlsxwriter
        output = BytesIO(); book = xlsxwriter.Workbook(output); book.add_worksheet().write('A1', 'Exact'); book.close(); data = output.getvalue()
    elif name.endswith('.odt'):
        from odf.opendocument import OpenDocumentText
        from odf.text import P
        text = OpenDocumentText(); text.text.addElement(P(text='Exact')); output = BytesIO(); text.save(output); data = output.getvalue()
    elif name.endswith('.ods'):
        from wiki.file_formats import write_ods
        data = write_ods('A,B\n1,2', 'Exact')
    else:
        data = b'# Markdown\n\nExact text.'
    client = Client(); client.force_login(workspace['admin'])
    response = client.post('/files/upload/', {'file':SimpleUploadedFile(name, data), 'path':'/'})
    assert response.status_code == 302, response.content[:200]
    doc = Document.objects.get(title=name)
    assert active_storage().read(doc.reference) == data
    page = client.get(f'/documents/{doc.pk}/')
    assert page.status_code == 200
    assert f'/documents/{doc.pk}/live/'.encode() in page.content
    assert client.get(f'/documents/{doc.pk}/live/').url.endswith('/office/')
    session = open_session(doc, workspace['admin'])
    assert client.get(url(session, workspace['admin'], True)).status_code == 200


def test_markdown_engine_save_and_binary_format_rejection(office, workspace):
    from django.core.exceptions import ValidationError
    from wiki.collabora import validate_file, eligible, editor_bytes
    from wiki.models import Document
    doc = save_document(workspace['admin'], 'Markdown', 'document', 'Notes', '# Initial', workspace['team'], file_format='md')
    session = open_session(doc, workspace['admin'])
    body = b'# Updated\n\nSaved in Collabora'
    response = Client().post(url(session, workspace['admin'], True), body,
                            content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT')
    assert response.status_code == 200
    assert active_storage().read(doc.reference) == body
    assert editor_bytes(doc, body) == body
    for name in ('Broken.pptx', 'Broken.doc', 'Broken.odt', 'Broken.rtf'):
        with pytest.raises(ValidationError):
            validate_file(Document(kind='file', title=name), b'bad file')
    for name in ('Drawing.drawio', 'Canvas.canvas', 'View.pdf', 'Image.png'):
        assert not eligible(Document(kind='file', title=name))


def test_office_heartbeat_rotates_external_version_and_keeps_late_edits(office, workspace, client):
    doc, session = office
    before = session.base_digest
    external = rich_docx()
    active_storage().write('ignored.docx', external, doc.reference)
    client.force_login(workspace['admin'])
    response = client.post(f'/documents/{doc.pk}/office/status/', {'session':str(session.pk)})
    assert response.status_code == 200, response.content
    assert response.json()['changed']
    session.refresh_from_db(); doc.refresh_from_db()
    assert session.superseded and session.base_digest == before
    assert doc.storage_digest != before
    fresh = open_session(doc, workspace['member'])
    assert fresh.pk != session.pk
    assert open_session(doc, workspace['admin']).pk == fresh.pk
    local = WordDocument(); local.add_paragraph('Local changes preserved after external update')
    buffer = BytesIO(); local.save(buffer); payload = buffer.getvalue()
    response = client.post(url(session, workspace['admin'], True), payload,
                           content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT')
    assert response.status_code == 409
    assert active_storage().read(doc.reference) == external
    copy = client.get(f'/documents/{doc.pk}/office/recovery/?session={session.pk}')
    assert copy.status_code == 200 and copy.content == payload
    assert client.get(url(fresh, workspace['member'], True)).content == external


def test_office_status_does_not_expose_another_authors_recovery(office, workspace, client):
    doc, session = office
    OfficeSession.objects.filter(pk=session.pk).update(pending=rich_docx(), pending_author=workspace['admin'])
    client.force_login(workspace['member'])
    response = client.post(f'/documents/{doc.pk}/office/status/', {'session':str(session.pk)})
    assert response.status_code == 200
    assert response.json()['recovery'] is False


@pytest.mark.parametrize('alignment,rtl,expected', [('start',False,'left'),('end',False,'right'),('start',True,'right'),('end',True,'left')])
def test_collabora_docx_logical_alignment_reopens_and_permissions_remain_editable(office, workspace, client, alignment, rtl, expected):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from wiki.file_formats import decode
    from wiki.native import unpack
    doc, session = office
    word = WordDocument(); para = word.add_paragraph('Text remains readable after Collabora save')
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'),alignment); para._p.get_or_add_pPr().append(jc)
    if rtl: para._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
    buffer = BytesIO(); word.save(buffer); payload = buffer.getvalue()
    response = client.post(url(session,workspace['admin'],True),payload,content_type='application/octet-stream',HTTP_X_WOPI_OVERRIDE='PUT')
    assert response.status_code == 200
    doc.refresh_from_db()
    assert active_storage().read(doc.reference) == payload
    native = unpack(decode(doc,payload))
    assert native['data']['content'][0]['attrs']['textAlign'] == expected
    client.force_login(workspace['admin'])
    assert client.get(f'/documents/{doc.pk}/').status_code == 200
    assert client.get(f'/documents/{doc.pk}/file-settings/').status_code == 200
    response = client.post(f'/documents/{doc.pk}/file-settings/', {'action':'access','group':doc.group_id,'owner.read':'on','owner.write':'on','group.read':'on'})
    assert response.status_code == 302
    doc.refresh_from_db()
    assert doc.policy['group'] == {'read':True,'write':False}

def test_trash_revokes_wopi_and_old_session_cannot_write_after_restore(office, workspace):
    from wiki.services import delete_document
    from wiki.models import TrashEntry
    from wiki.trash import restore
    doc, session = office
    client = Client()
    info_url = url(session, workspace['member'])
    contents_url = url(session, workspace['member'], contents=True)
    original = active_storage().read(doc.reference)
    delete_document(workspace['admin'], doc)
    assert client.get(info_url).status_code == 403
    assert client.get(contents_url).status_code == 403
    assert client.post(contents_url, data=rich_docx(), content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT').status_code == 403
    restore(workspace['admin'], TrashEntry.objects.get())
    response = client.post(contents_url, data=rich_docx(), content_type='application/octet-stream', HTTP_X_WOPI_OVERRIDE='PUT')
    assert response.status_code == 409
    assert active_storage().read(doc.reference) == original
