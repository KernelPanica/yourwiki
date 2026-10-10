"""Document toolbar formatting must survive portable Word files and previews."""
from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from wiki.docx import to_docx, to_native
from wiki.native import pack, unpack, rich_html, safe_style


def test_docs_gui_upgrade_word_run_and_paragraph_roundtrip():
    text = {'type': 'text', 'text': 'Formatted title', 'marks': [
        {'type': 'bold'}, {'type': 'italic'}, {'type': 'underline'}, {'type': 'strike'},
        {'type': 'textStyle', 'attrs': {'fontFamily': 'Times New Roman', 'fontSize': '18pt', 'color': '#123456'}},
        {'type': 'highlight', 'attrs': {'color': '#abcdef'}},
        {'type': 'link', 'attrs': {'href': 'https://example.com'}}]}
    native = pack('document', {'type': 'doc', 'content': [
        {'type': 'heading', 'attrs': {'level': 6, 'textAlign': 'center', 'lineHeight': '2'}, 'content': [text]}]})
    data = to_docx(native)
    paragraph = Document(BytesIO(data)).paragraphs[0]
    assert paragraph.style.name == 'Heading 6'
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.paragraph_format.line_spacing == 2
    imported = unpack(to_native(BytesIO(data)))['data']['content'][0]
    assert imported['attrs'] == {'level': 6, 'textAlign': 'center', 'lineHeight': '2'}
    marks = {mark['type']: mark.get('attrs') for mark in imported['content'][0]['marks']}
    assert {'bold', 'italic', 'underline', 'strike', 'link', 'highlight'} <= marks.keys()
    assert marks['textStyle'] == {'fontFamily': 'Times New Roman', 'fontSize': '18pt', 'color': '#123456'}
    assert marks['highlight']['color'].lower() == '#abcdef'
    assert marks['link']['href'] == 'https://example.com'


def test_docs_gui_upgrade_legacy_px_size_and_hard_break():
    data = to_docx(pack('document', {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [
        {'type': 'text', 'text': 'Before', 'marks': [{'type': 'textStyle', 'attrs': {'fontSize': '16px'}}]},
        {'type': 'hardBreak'}, {'type': 'text', 'text': 'After'}]}]}))
    paragraph = Document(BytesIO(data)).paragraphs[0]
    assert paragraph.runs[0].font.size.pt == 12
    imported = unpack(to_native(BytesIO(data)))['data']['content'][0]
    assert any(node['type'] == 'hardBreak' for node in imported['content'])


def test_docs_gui_upgrade_list_indentation_and_table_merge_roundtrip():
    def para(text):
        return {'type': 'paragraph', 'content': [{'type': 'text', 'text': text}]}
    listing = {'type': 'bulletList', 'content': [{'type': 'listItem', 'content': [para('Parent'),
        {'type': 'bulletList', 'content': [{'type': 'listItem', 'content': [para('Nested')]}]}]}]}
    table = {'type': 'table', 'content': [
        {'type': 'tableRow', 'content': [
            {'type': 'tableCell', 'attrs': {'rowspan': 2, 'backgroundColor': '#ffddaa'}, 'content': [para('Merged')]},
            {'type': 'tableCell', 'content': [para('Top')]}]},
        {'type': 'tableRow', 'content': [{'type': 'tableCell', 'content': [para('Bottom')]}]}]}
    data = to_docx(pack('document', {'type': 'doc', 'content': [listing, table]}))
    imported = unpack(to_native(BytesIO(data)))['data']['content']
    assert imported[0]['content'][0]['content'][1]['type'] == 'bulletList'
    merged = imported[1]['content'][0]['content'][0]
    assert merged['attrs']['rowspan'] == 2
    assert merged['attrs']['backgroundColor'].lower() == '#ffddaa'
    assert len(imported[1]['content'][1]['content']) == 1
    assert Document(BytesIO(data)).tables[0].cell(1, 1).text == 'Bottom'


def test_docs_gui_upgrade_preview_safe_fonts_spacing_and_cell_fill():
    assert 'font-family:Times New Roman' in safe_style({'fontFamily': 'Times New Roman'})
    assert 'font-size:10.5pt' in safe_style({'fontSize': '10.5pt'})
    assert safe_style({'fontSize': '90pt'}) == ''
    assert 'line-height:1.15' in safe_style({'lineHeight': '1.15'})
    assert safe_style({'fontFamily': 'Arial;position:fixed', 'lineHeight': 'url(evil)'}) == ''
    preview = rich_html({'type': 'tableCell', 'attrs': {'backgroundColor': '#abcdef'}, 'content': []})
    assert 'background-color:#abcdef' in preview


def test_docs_gui_upgrade_word_table_header_column():
    data = to_docx(pack('document', {'type': 'doc', 'content': [{'type': 'table', 'content': [
        {'type': 'tableRow', 'content': [
            {'type': 'tableHeader', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': 'Label'}]}]},
            {'type': 'tableCell', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': 'Value'}]}]}]}]}]}))
    assert Document(BytesIO(data)).tables[0].cell(0, 0).paragraphs[0].runs[0].bold
    row = unpack(to_native(BytesIO(data)))['data']['content'][0]['content'][0]
    assert [cell['type'] for cell in row['content']] == ['tableHeader', 'tableCell']
