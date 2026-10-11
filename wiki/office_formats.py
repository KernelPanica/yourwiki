"""Office formats advertised by the pinned Collabora CODE discovery."""
import mimetypes
from pathlib import PurePosixPath

TEXT = {'md', 'txt', 'csv', 'tsv', 'dif', 'slk'}
OLE = {'doc', 'xls', 'ppt', 'xla'}
OPENXML = {
    **dict.fromkeys(('docx', 'docm', 'dotx', 'dotm'), 'word/document.xml'),
    **dict.fromkeys(('xlsx', 'xlsm', 'xltx', 'xltm'), 'xl/workbook.xml'),
    'xlsb': 'xl/workbook.bin',
    **dict.fromkeys(('pptx', 'pptm', 'potx', 'potm', 'ppsx'), 'ppt/presentation.xml'),
}
ODF = dict(zip(('odt', 'ods', 'odp', 'odg', 'odm', 'ott', 'ots', 'otp', 'oth'),
    ('text', 'spreadsheet', 'presentation', 'graphics', 'text-master', 'text-template', 'spreadsheet-template', 'presentation-template', 'text-web')))
FLAT = {'fodt', 'fods', 'fodp', 'fodg'}
SUPPORTED = TEXT | OLE | set(OPENXML) | set(ODF) | FLAT | {'rtf', 'dbf'}


def extension(document):
    return PurePosixPath(document.title).suffix.lower().lstrip('.') if document.kind == 'file' else document.file_format


def mime(extension):
    if extension in ODF:
        return 'application/vnd.oasis.opendocument.' + ODF[extension]
    return mimetypes.guess_type('file.' + extension)[0] or 'application/octet-stream'


def presentation(title, content=''):
    """Create a standard PPTX; subsequent editing saves the engine's exact bytes."""
    from pptx import Presentation
    from io import BytesIO
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = content
    output = BytesIO()
    deck.save(output)
    return output.getvalue()
