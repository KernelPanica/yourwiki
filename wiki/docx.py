"""DOCX import and export helpers."""
from io import BytesIO
import re
from .native import pack


def _text(paragraph):
    from docx.oxml.ns import qn
    content = []
    # iter_inner_content includes hyperlinks while preserving run order.
    for child in paragraph.iter_inner_content():
        runs = child.runs if hasattr(child, 'runs') else [child]
        for run in runs:
            marks = []
            for kind, active in [('bold', run.bold), ('italic', run.italic), ('underline', run.underline), ('strike', run.font.strike)]:
                if active:
                    marks.append({'type': kind})
            attrs = {}
            if run.font.name:
                attrs['fontFamily'] = run.font.name
            if run.font.size:
                attrs['fontSize'] = f'{run.font.size.pt:g}pt'
            if run.font.color.rgb:
                attrs['color'] = '#' + str(run.font.color.rgb)
            if attrs:
                marks.append({'type': 'textStyle', 'attrs': attrs})
            shading = run._element.find(qn('w:rPr'))
            shading = shading.find(qn('w:shd')) if shading is not None else None
            fill = shading.get(qn('w:fill')) if shading is not None else None
            if fill and re.fullmatch(r'[0-9a-fA-F]{6}', fill):
                marks.append({'type': 'highlight', 'attrs': {'color': '#' + fill}})
            elif run.font.highlight_color:
                colors = {1: '#000000', 2: '#0000ff', 3: '#00ffff', 4: '#00ff00', 5: '#ff00ff', 6: '#ff0000', 7: '#ffff00', 8: '#ffffff', 9: '#000080', 10: '#008080', 11: '#008000', 12: '#800080', 13: '#800000', 14: '#808000', 15: '#808080', 16: '#c0c0c0'}
                color = colors.get(int(run.font.highlight_color))
                if color:
                    marks.append({'type': 'highlight', 'attrs': {'color': color}})
            if hasattr(child, 'url') and child.url.startswith(('http:', 'https:', 'mailto:')):
                marks.append({'type': 'link', 'attrs': {'href': child.url}})
            parts = run.text.split('\n')
            for index, text in enumerate(parts):
                if index:
                    content.append({'type': 'hardBreak'})
                if text:
                    content.append({'type': 'text', 'text': text, 'marks': marks})
    if not content and paragraph.text:
        content.append({'type': 'text', 'text': paragraph.text, 'marks': []})
    return content


def _paragraph(paragraph):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    style = paragraph.style.name if paragraph.style else ''
    attrs = {}
    alignments = {WD_ALIGN_PARAGRAPH.LEFT: 'left', WD_ALIGN_PARAGRAPH.CENTER: 'center', WD_ALIGN_PARAGRAPH.RIGHT: 'right', WD_ALIGN_PARAGRAPH.JUSTIFY: 'justify'}
    if paragraph.alignment in alignments:
        attrs['textAlign'] = alignments[paragraph.alignment]
    spacing = paragraph.paragraph_format.line_spacing
    if isinstance(spacing, float) and spacing in (1, 1.15, 1.5, 2):
        attrs['lineHeight'] = f'{spacing:g}'
    kind = 'paragraph'
    if style.startswith('Heading '):
        try:
            kind, attrs['level'] = 'heading', max(1, min(6, int(style.rsplit(' ', 1)[1])))
        except ValueError:
            pass
    node = {'type': kind, 'content': _text(paragraph)}
    if attrs:
        node['attrs'] = attrs
    return node


def _blocks(source):
    from docx.table import Table
    result, list_stack = [], []
    for block in source.iter_inner_content():
        if isinstance(block, Table):
            result.append(_table(block))
            list_stack = []
            continue
        style = block.style.name if block.style else ''
        node = _paragraph(block)
        if style.startswith(('List Bullet', 'List Number')):
            kind = 'bulletList' if style.startswith('List Bullet') else 'orderedList'
            match = re.search(r' (\d+)$', style)
            level = max(0, min(2, int(match[1])-1)) if match else 0
            level = min(level, len(list_stack))
            list_stack = list_stack[:level+1]
            if len(list_stack) <= level or list_stack[level]['type'] != kind:
                container = result if level == 0 else list_stack[level-1]['content'][-1]['content']
                listing = {'type': kind, 'content': []}
                container.append(listing)
                list_stack = list_stack[:level] + [listing]
            list_stack[level]['content'].append({'type': 'listItem', 'content': [node]})
        elif style == 'Quote':
            result.append({'type': 'blockquote', 'content': [node]})
            list_stack = []
        else:
            result.append(node)
            list_stack = []
    return result


def _table(table):
    from docx.oxml.ns import qn
    rows, seen = [], set()
    for row in table.rows:
        cells = []
        header = row._tr.trPr is not None and row._tr.trPr.find(qn('w:tblHeader')) is not None
        for cell in row.cells:
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            attrs = {'colspan': cell._tc.grid_span}
            # python-docx resolves merged cells to the origin; count rows it spans.
            attrs['rowspan'] = sum(any(other._tc is cell._tc for other in candidate.cells) for candidate in table.rows)
            shading = cell._tc.tcPr.find(qn('w:shd'))
            fill = shading.get(qn('w:fill')) if shading is not None else None
            if fill and re.fullmatch(r'[0-9a-fA-F]{6}', fill):
                attrs['backgroundColor'] = '#' + fill
            conditional = cell._tc.tcPr.find(qn('w:cnfStyle'))
            cell_header = header or conditional is not None and conditional.get(qn('w:firstColumn')) == '1'
            cells.append({'type': 'tableHeader' if cell_header else 'tableCell', 'attrs': attrs, 'content': _blocks(cell) or [{'type': 'paragraph'}]})
        rows.append({'type': 'tableRow', 'content': cells})
    return {'type': 'table', 'content': rows}


def to_native(data):
    from docx import Document
    source = Document(data)
    return pack('document', {'type': 'doc', 'content': _blocks(source) or [{'type': 'paragraph'}]})


def to_docx(native_text):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from .native import unpack
    native = unpack(native_text) or {'data': {'content': []}}
    output = Document()

    def shade(properties, color):
        if isinstance(color, str) and re.fullmatch(r'#[0-9a-fA-F]{6}', color):
            element = OxmlElement('w:shd')
            element.set(qn('w:fill'), color[1:])
            properties.append(element)

    def blocks(parent, nodes, list_style=None, depth=0):
        for node in nodes:
            kind, children, attrs = node.get('type'), node.get('content', []), node.get('attrs') or {}
            if kind == 'table':
                rows = children
                # Rowspans occupy columns even when no node exists in the next row.
                occupied, placements, columns = set(), [], 0
                for r, row in enumerate(rows):
                    c = 0
                    for cell in row.get('content', []):
                        while (r, c) in occupied:
                            c += 1
                        ca = cell.get('attrs') or {}
                        width, height = max(1, min(100, int(ca.get('colspan', 1)))), max(1, min(100, int(ca.get('rowspan', 1))))
                        height = min(height, len(rows)-r)
                        placements.append((r, c, width, height, cell))
                        occupied.update((rr, cc) for rr in range(r, r+height) for cc in range(c, c+width))
                        c += width
                        columns = max(columns, c)
                if not rows or not columns:
                    continue
                table = parent.add_table(rows=len(rows), cols=columns)
                table.style = 'Table Grid'
                for r, row in enumerate(rows):
                    if row.get('content') and all(cell.get('type') == 'tableHeader' for cell in row['content']):
                        table.rows[r]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for r, c, width, height, cell in placements:
                    target = table.cell(r, c)
                    if width > 1 or height > 1:
                        target = target.merge(table.cell(r+height-1, c+width-1))
                    shade(target._tc.get_or_add_tcPr(), (cell.get('attrs') or {}).get('backgroundColor'))
                    if cell.get('type') == 'tableHeader':
                        conditional = OxmlElement('w:cnfStyle')
                        conditional.set(qn('w:firstColumn'), '1')
                        target._tc.get_or_add_tcPr().append(conditional)
                    blocks(target, cell.get('content', []))
                    if cell.get('type') == 'tableHeader':
                        for paragraph in target.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
                    if len(target.paragraphs) > 1:
                        empty = target.paragraphs[0]._element
                        empty.getparent().remove(empty)
            elif kind in ('bulletList', 'orderedList'):
                level = min(depth, 2)
                style = 'List Bullet' if kind == 'bulletList' else 'List Number'
                blocks(parent, children, style + (f' {level+1}' if level else ''), depth+1)
            elif kind in ('listItem', 'blockquote'):
                blocks(parent, children, list_style or ('Quote' if kind == 'blockquote' else None), depth)
            elif kind in ('paragraph', 'heading', 'codeBlock'):
                paragraph = parent.add_paragraph(style=list_style)
                if kind == 'heading':
                    paragraph.style = f"Heading {max(1, min(6, int(attrs.get('level', 1))))}"
                alignment = attrs.get('textAlign')
                if alignment in ('left', 'center', 'right', 'justify'):
                    paragraph.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[alignment]
                if str(attrs.get('lineHeight')) in ('1', '1.15', '1.5', '2'):
                    paragraph.paragraph_format.line_spacing = float(attrs['lineHeight'])
                for child in children:
                    run = paragraph.add_run(child.get('text', ''))
                    if child.get('type') == 'hardBreak':
                        run.add_break()
                    marks = {mark.get('type'): mark.get('attrs') or {} for mark in child.get('marks', [])}
                    run.bold, run.italic, run.underline = 'bold' in marks, 'italic' in marks, 'underline' in marks
                    run.font.strike = 'strike' in marks
                    if kind == 'codeBlock' or 'code' in marks:
                        run.font.name = 'Courier New'
                    text = marks.get('textStyle', {})
                    if text.get('fontFamily'):
                        run.font.name = text['fontFamily']
                    size = re.fullmatch(r'(\d+(?:\.\d+)?)(px|pt)', str(text.get('fontSize', '')))
                    if size and 1 <= float(size[1]) <= 200:
                        run.font.size = Pt(float(size[1]) * (.75 if size[2] == 'px' else 1))
                    if re.fullmatch(r'#[0-9a-fA-F]{6}', str(text.get('color', ''))):
                        run.font.color.rgb = RGBColor.from_string(text['color'][1:])
                    shade(run._element.get_or_add_rPr(), marks.get('highlight', {}).get('color'))
                    href = marks.get('link', {}).get('href', '')
                    if isinstance(href, str) and href.startswith(('http:', 'https:', 'mailto:')):
                        from docx.opc.constants import RELATIONSHIP_TYPE as RT
                        link = OxmlElement('w:hyperlink')
                        link.set(qn('r:id'), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
                        paragraph._p.remove(run._r)
                        link.append(run._r)
                        paragraph._p.append(link)
            elif children:
                blocks(parent, children, list_style, depth)
    blocks(output, native['data'].get('content', []))
    stream = BytesIO()
    output.save(stream)
    return stream.getvalue()
